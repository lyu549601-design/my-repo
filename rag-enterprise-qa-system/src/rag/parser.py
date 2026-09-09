"""文档解析与切片器：Word/PDF/PPT/Excel 表格行级切片。"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

from .models import DocumentChunk


def estimate_tokens(text: str) -> int:
    """Token 估算：中文按字符数近似，英文按 4 字符 1 Token。"""
    if not text:
        return 0
    cjk_count = len(re.findall(r"[\u4e00-\u9fff]", text))
    other_count = len(re.sub(r"[\u4e00-\u9fff]", "", text))
    return cjk_count + max(1, other_count // 4)


def table_to_markdown_rows(headers: list[str], rows: list[list[str]]) -> list[str]:
    """将表格转成带表头的 Markdown 行级文本。"""
    if not headers:
        return []
    header_line = "| " + " | ".join(str(item) for item in headers) + " |"
    output: list[str] = []
    for row in rows:
        cells = [str(cell) if cell is not None else "" for cell in row]
        output.append(header_line + "\n" + "| " + " | ".join(cells) + " |")
    return output


def split_text_into_chunks(text: str, max_tokens: int = 500) -> list[str]:
    """按句边界切分，单块不超过 max_tokens。"""
    if estimate_tokens(text) <= max_tokens:
        return [text] if text.strip() else []

    sentences = re.split(r"(?<=[。！？；\n])", text)
    chunks: list[str] = []
    current = ""

    for sentence in sentences:
        if not sentence.strip():
            continue
        if estimate_tokens(current + sentence) <= max_tokens:
            current += sentence
            continue

        if current:
            chunks.append(current.strip())
            current = ""

        remaining = sentence
        while estimate_tokens(remaining) > max_tokens:
            cut = min(max_tokens, max(1, len(remaining) - 1))
            chunks.append(remaining[:cut].strip())
            remaining = remaining[cut:]
        current = remaining

    if current.strip():
        chunks.append(current.strip())
    return [chunk for chunk in chunks if chunk]


class DocumentParser:
    """按扩展名解析文档并生成切片。"""

    def __init__(self, max_chunk_tokens: int = 500):
        self.max_chunk_tokens = max_chunk_tokens

    async def parse_file(
        self,
        path: str,
        document_id: str,
        title: str,
    ) -> list[DocumentChunk]:
        file_path = Path(path)
        suffix = file_path.suffix.lower()
        if suffix in {".docx", ".doc"}:
            return self._parse_docx(path, document_id, title)
        if suffix == ".pdf":
            return self._parse_pdf(path, document_id, title)
        if suffix in {".pptx", ".ppt"}:
            return self._parse_pptx(path, document_id, title)
        if suffix in {".xlsx", ".xls"}:
            return self._parse_xlsx(path, document_id, title)
        raise ValueError(f"不支持的文档格式: {suffix}")

    def _parse_docx(self, path: str, document_id: str, title: str) -> list[DocumentChunk]:
        from docx import Document

        doc = Document(path)
        chunks: list[DocumentChunk] = []
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text)
        for table in doc.tables:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in table.rows[1:]
            ]
            text += "\n" + "\n".join(table_to_markdown_rows(headers, rows))
        self._append_text_chunks(chunks, text, document_id, title, page=1)
        return chunks

    def _parse_pdf(self, path: str, document_id: str, title: str) -> list[DocumentChunk]:
        import pdfplumber

        chunks: list[DocumentChunk] = []
        with pdfplumber.open(path) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                for table in page.extract_tables() or []:
                    if not table:
                        continue
                    headers = [str(cell or "") for cell in table[0]]
                    rows = [
                        [str(cell or "") for cell in row]
                        for row in table[1:]
                    ]
                    text += "\n" + "\n".join(table_to_markdown_rows(headers, rows))
                self._append_text_chunks(chunks, text, document_id, title, page=page_index)
        return chunks

    def _parse_pptx(self, path: str, document_id: str, title: str) -> list[DocumentChunk]:
        from pptx import Presentation

        presentation = Presentation(path)
        chunks: list[DocumentChunk] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            text = "\n".join(
                shape.text
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text
            )
            self._append_text_chunks(chunks, text, document_id, title, page=slide_index)
        return chunks

    def _parse_xlsx(self, path: str, document_id: str, title: str) -> list[DocumentChunk]:
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=True, read_only=True)
        chunks: list[DocumentChunk] = []
        for sheet_index, sheet_name in enumerate(workbook.sheetnames, start=1):
            sheet = workbook[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(cell or "") for cell in rows[0]]
            data_rows = rows[1:]
            for line in table_to_markdown_rows(headers, data_rows):
                chunks.append(
                    self._make_chunk(
                        document_id=document_id,
                        title=title,
                        page=sheet_index,
                        chapter=sheet_name,
                        content=line,
                    )
                )
        return chunks

    def _append_text_chunks(
        self,
        chunks: list[DocumentChunk],
        text: str,
        document_id: str,
        title: str,
        page: int,
    ) -> None:
        for part in split_text_into_chunks(text, self.max_chunk_tokens):
            chunks.append(
                self._make_chunk(
                    document_id=document_id,
                    title=title,
                    page=page,
                    content=part,
                )
            )

    def _make_chunk(
        self,
        document_id: str,
        title: str,
        page: int,
        content: str,
        chapter: str = "",
    ) -> DocumentChunk:
        digest = hashlib.sha256(content.encode("utf-8")).hexdigest()[:12]
        return DocumentChunk(
            id=f"{document_id}_p{page}_{digest}",
            document_id=document_id,
            page=page,
            chapter=chapter,
            content=content,
            metadata={"document_title": title},
        )
